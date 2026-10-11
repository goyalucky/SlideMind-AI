import os
from pptx import Presentation

class SlideGenerator:
    def __init__(self, output_dir=None):
        if output_dir is None:
            # Resolves path to the project-level output/ folder
            self.output_dir = os.path.abspath(
                os.path.join(os.path.dirname(__file__), "../../output")
            )
        else:
            self.output_dir = output_dir

        os.makedirs(self.output_dir, exist_ok=True)

    def create_slide_deck(self, title: str, bullet_points: list):
        prs = Presentation()

        # Title slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title_placeholder = slide.shapes.title
        title_placeholder.text = title

        # Content slides (up to 5 bullets per slide)
        for i in range(0, len(bullet_points), 5):
            bullet_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_slide_layout)
            
            title_shape = slide.shapes.title
            content_shape = slide.placeholders[1]
            
            title_shape.text = f"{title} (Part {i//5 + 1})"
            tf = content_shape.text_frame
            
            batch = bullet_points[i:i + 5]
            for idx, bullet in enumerate(batch):
                if idx == 0:
                    tf.text = bullet
                else:
                    p = tf.add_paragraph()
                    p.text = bullet

        pptx_path = os.path.join(self.output_dir, "generated_presentation.pptx")
        prs.save(pptx_path)
        print(f"\nPowerPoint presentation saved to: {pptx_path}\n")


if __name__ == "__main__":
    sample_title = "AI-Based Slide Generation"
    sample_points = [
        "AI automates repetitive tasks and analyzes vast data.",
        "Continuous learning and project practice build key skills.",
        "Clear slides present concise, actionable takeaways.",
        "Integrating LLMs streamlines presentation workflows.",
        "Scalable architectures allow automated slide decks from raw text."
    ]

    slide_generator = SlideGenerator()
    slide_generator.create_slide_deck(sample_title, sample_points)